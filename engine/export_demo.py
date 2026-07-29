import random
import datetime
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from scheduler import run_month, SHIFT_WINDOWS, STRETCH_MIN_CONSECUTIVE

random.seed(1)
dates, drivers, stretches, tasks_by_shift, uncovered_tasks = run_month()

OUT = "Scheduling_Engine_Demo.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Demo Schedule"

ARIAL = "Arial"
header_font = Font(name=ARIAL, bold=True, size=10, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="1F3864")
green = Font(name=ARIAL, size=9, color="008000", bold=True)
orange_fill = PatternFill("solid", fgColor="F4B183")
grey = Font(name=ARIAL, size=9, color="808080")

ws["A1"] = "Driver"
ws["B1"] = "Shift type per stretch"
for i, d in enumerate(dates):
    c = ws.cell(row=1, column=3 + i, value=d)
    c.number_format = "dd-mmm"
for coord in ("A1", "B1"):
    ws[coord].font = header_font
    ws[coord].fill = header_fill
for i in range(len(dates)):
    cell = ws.cell(row=1, column=3 + i)
    cell.font = header_font
    cell.fill = header_fill

by_driver = defaultdict(list)
for s in stretches:
    by_driver[s.driver].append(s)

row = 2
for name, drv_stretches in by_driver.items():
    drv_stretches.sort(key=lambda s: s.day_indexes[0])
    ws.cell(row=row, column=1, value=name).font = Font(name=ARIAL, size=9)
    shift_seq = " -> ".join(s.shift_type for s in drv_stretches)
    ws.cell(row=row, column=2, value=shift_seq).font = Font(name=ARIAL, size=8)
    for s in drv_stretches:
        for day_idx, val in s.assignments.items():
            cell = ws.cell(row=row, column=3 + day_idx, value=val)
            if val == "RESERVE":
                cell.fill = orange_fill
                cell.font = Font(name=ARIAL, size=9, bold=True)
            else:
                cell.font = green
    row += 1

ws.freeze_panes = "C2"
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 30

# ---- compliance summary sheet ----
ws2 = wb.create_sheet("Compliance Summary")
ws2["A1"] = "Metric"
ws2["B1"] = "Value"
ws2["A1"].font = header_font
ws2["B1"].font = header_font
ws2["A1"].fill = header_fill
ws2["B1"].fill = header_fill

total_stretches = len(stretches)
ratio_exact = 0
ratio_off = defaultdict(int)
violations = 0
for name, drv_stretches in by_driver.items():
    drv_stretches.sort(key=lambda s: s.day_indexes[0])
    for i, s in enumerate(drv_stretches):
        actual = sum(1 for v in s.assignments.values() if v != "RESERVE")
        diff = actual - s.target_trips
        if diff == 0:
            ratio_exact += 1
        else:
            ratio_off[diff] += 1
        if i > 0:
            prev = drv_stretches[i - 1]
            if prev.shift_type != s.shift_type:
                run = 1
                j = i - 1
                while j > 0 and drv_stretches[j - 1].shift_type == prev.shift_type:
                    run += 1
                    j -= 1
                if run < STRETCH_MIN_CONSECUTIVE:
                    violations += 1

rows = [
    ("Total work stretches detected", total_stretches),
    ("Stretches hitting exact trip target", f"{ratio_exact} ({100*ratio_exact/total_stretches:.1f}%)"),
    ("Stretches below target (fewer trips than aimed for)", sum(v for k, v in ratio_off.items() if k < 0)),
    ("Stretches above target", sum(v for k, v in ratio_off.items() if k > 0)),
    ("Shift-stickiness rule violations (should always be 0)", violations),
    ("Task-days genuinely short a driver (real shift-type shortage)", len(uncovered_tasks)),
    ("", ""),
    ("Task count by shift type:", ""),
]
for label, _, _ in SHIFT_WINDOWS:
    rows.append((f"  {label}", len(tasks_by_shift.get(label, []))))

for i, (label, val) in enumerate(rows, start=2):
    ws2.cell(row=i, column=1, value=label).font = Font(name=ARIAL, size=10, bold=label.strip().endswith(":"))
    ws2.cell(row=i, column=2, value=val).font = Font(name=ARIAL, size=10)

ws2.column_dimensions["A"].width = 55
ws2.column_dimensions["B"].width = 20

wb.save(OUT)
print("saved", OUT)
