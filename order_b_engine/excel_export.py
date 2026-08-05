"""Write a list of Duty objects to an Excel task schedule, styled like the reference sheet:
bordered trip-number box, route-family color bar underneath, times rotated on either side,
turnaround-station letter badge, and a RESERVE block for reserve duties.

Note: the turnaround badge here is a plain colored cell with the letter (not a floating circle
image) -- an earlier prototype used an embedded circle image, but it doesn't render in common
quick-look/mobile previewers, so this generator standardizes on the colored-cell version, which
is confirmed to render everywhere. See README.md for that back-and-forth if you want the true
circle restored for desktop-Excel-only use.
"""
from __future__ import annotations

import datetime as dt
import re
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import Duty, Role
from .trip_codes import route_color, STATION_LETTER

FONT_NAME = "Arial"

_COLOR_HEX = {"blue": "4472C4", "green": "70AD47", "red": "C00000"}
HEADER_FILL = "1F2A44"
STRIPE_FILL = "D9EAD3"
WHITE = "FFFFFF"
RESERVE_FILL = "BFBFBF"
OT_FILL = "FFC000"

thin = Side(style="thin", color="000000")
medium = Side(style="medium", color="000000")
box_border = Border(left=medium, right=medium, top=medium, bottom=medium)
plain_border = Border(left=thin, right=thin, top=thin, bottom=thin)

header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
title_font = Font(name=FONT_NAME, bold=True, size=14)
sub_font = Font(name=FONT_NAME, italic=True, size=9, color="555555")
normal_font = Font(name=FONT_NAME, size=10)
bold_font = Font(name=FONT_NAME, bold=True, size=10)
red_font = Font(name=FONT_NAME, bold=True, size=10, color="C00000")
trip_font = Font(name=FONT_NAME, bold=True, size=12, color="000000")
time_font = Font(name=FONT_NAME, bold=True, size=8, color="000000")
driver_font = Font(name=FONT_NAME, bold=True, size=9, color="000000")
reserve_font = Font(name=FONT_NAME, bold=True, size=12, color="FFFFFF")

vertical_up = Alignment(text_rotation=90, horizontal="center", vertical="center")
center = Alignment(horizontal="center", vertical="center")
wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Columns: A Task Code | B Driver Info | C Sign In | D Origin |
# E leg1-dep | F leg1-box | G leg1-arr | H turn | I leg2-dep | J leg2-box | K leg2-arr |
# L Destination | M Sign Out | N Total Duty | O OT
_HEADER = {
    "A": "Task Code", "B": "Driver Info (ID / Name / Phone)", "C": "Sign In", "D": "Origin",
    "E": "Leg 1", "H": "Turn", "I": "Leg 2",
    "L": "Destination", "M": "Sign Out", "N": "Total Duty", "O": "OT",
}


def _time_str(t: dt.time) -> str:
    return t.strftime("%H:%M")


def _write_legend(ws, row: int) -> int:
    ws.cell(row=row, column=1, value="Legend:").font = bold_font
    items = [
        (_COLOR_HEX["blue"], "MAK ⇄ MAD trips (00 / 01 / 03)"),
        (_COLOR_HEX["green"], "MAD ⇄ KAIA trips (07 / 08)"),
        (_COLOR_HEX["red"], "Shuttle trips MAK ⇄ KAIA (05)"),
    ]
    col = 2
    for color, label in items:
        c = ws.cell(row=row, column=col, value="  ")
        c.fill = PatternFill("solid", fgColor=color)
        c.border = plain_border
        lbl = ws.cell(row=row, column=col + 1, value=label)
        lbl.font = normal_font
        ws.merge_cells(start_row=row, start_column=col + 1, end_row=row, end_column=col + 3)
        col += 4
    ws.cell(row=row + 1, column=1, value="Station letters:").font = bold_font
    ws.cell(row=row + 1, column=2, value="L = MAK   M = MAD   A = KAIA   K = KAEC").font = normal_font
    ws.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=6)
    return row + 3


def _write_header_row(ws, row: int):
    for col_letter, h in _HEADER.items():
        c = ws[f"{col_letter}{row}"]
        c.value = h
        c.font = header_font
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = wrap_center
        c.border = plain_border
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)   # E:G Leg1
    ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=11)  # I:K Leg2
    for blank_col in ["F", "G", "J", "K"]:
        ws[f"{blank_col}{row}"].fill = PatternFill("solid", fgColor=HEADER_FILL)
        ws[f"{blank_col}{row}"].border = plain_border
    ws.row_dimensions[row].height = 26


def _merge_v(ws, col: int, main_row: int, bar_row: int):
    ws.merge_cells(start_row=main_row, start_column=col, end_row=bar_row, end_column=col)


def _write_duty_row(ws, duty: Duty, main_row: int) -> None:
    bar_row = main_row + 1

    ws[f"A{main_row}"] = duty.task_code
    ws[f"A{main_row}"].font = normal_font
    ws[f"A{main_row}"].alignment = center
    ws[f"A{main_row}"].border = plain_border
    _merge_v(ws, 1, main_row, bar_row)

    driver = duty.driver
    ws[f"B{main_row}"] = f"ID: {driver.driver_id}\n{driver.name}\nTel: {driver.phone}"
    ws[f"B{main_row}"].font = driver_font
    ws[f"B{main_row}"].alignment = wrap_center
    ws[f"B{main_row}"].border = plain_border
    _merge_v(ws, 2, main_row, bar_row)

    ws[f"C{main_row}"] = duty.sign_in
    ws[f"C{main_row}"].number_format = "hh:mm"
    ws[f"C{main_row}"].font = red_font
    ws[f"C{main_row}"].alignment = center
    ws[f"C{main_row}"].border = plain_border
    _merge_v(ws, 3, main_row, bar_row)

    ws[f"M{main_row}"] = duty.sign_out
    ws[f"M{main_row}"].number_format = "hh:mm"
    ws[f"M{main_row}"].font = red_font
    ws[f"M{main_row}"].alignment = center
    ws[f"M{main_row}"].border = plain_border
    _merge_v(ws, 13, main_row, bar_row)

    ws[f"D{main_row}"] = duty.origin
    ws[f"L{main_row}"] = duty.destination
    for col_letter in ("D", "L"):
        c = ws[f"{col_letter}{main_row}"]
        c.font = normal_font
        c.alignment = center
        c.border = plain_border
        _merge_v(ws, c.column, main_row, bar_row)

    for col_letter in ["E", "F", "G", "H", "I", "J", "K"]:
        for rr in (main_row, bar_row):
            ws[f"{col_letter}{rr}"].fill = PatternFill("solid", fgColor=STRIPE_FILL)

    if duty.is_reserve:
        ws.merge_cells(start_row=main_row, start_column=5, end_row=bar_row, end_column=11)
        rc = ws[f"E{main_row}"]
        rc.value = "RESERVE"
        rc.font = reserve_font
        rc.alignment = center
        rc.fill = PatternFill("solid", fgColor=RESERVE_FILL)
        rc.border = box_border
    else:
        leg1, leg2 = duty.legs[0], duty.legs[1]
        color = _COLOR_HEX[route_color(leg1.trip.prefix)]

        ws[f"E{main_row}"] = leg1.trip.dep_time
        ws[f"E{main_row}"].number_format = "hh:mm"
        ws[f"E{main_row}"].font = time_font
        ws[f"E{main_row}"].alignment = vertical_up

        ws[f"F{main_row}"] = leg1.trip.trip_no
        ws[f"F{main_row}"].font = trip_font
        ws[f"F{main_row}"].alignment = center
        ws[f"F{main_row}"].border = box_border
        ws[f"F{main_row}"].fill = PatternFill("solid", fgColor=WHITE)
        ws[f"F{bar_row}"].fill = PatternFill("solid", fgColor=color)
        ws[f"F{bar_row}"].border = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws[f"G{main_row}"] = leg1.trip.arr_time
        ws[f"G{main_row}"].number_format = "hh:mm"
        ws[f"G{main_row}"].font = time_font
        ws[f"G{main_row}"].alignment = vertical_up

        turn_letter = STATION_LETTER.get(leg1.trip.destination, "?")
        _merge_v(ws, 8, main_row, bar_row)
        tc = ws[f"H{main_row}"]
        tc.value = turn_letter
        tc.font = Font(name=FONT_NAME, bold=True, size=12, color="FFFFFF")
        tc.alignment = center
        tc.fill = PatternFill("solid", fgColor=color)
        tc.border = plain_border

        color2 = _COLOR_HEX[route_color(leg2.trip.prefix)]
        ws[f"I{main_row}"] = leg2.trip.dep_time
        ws[f"I{main_row}"].number_format = "hh:mm"
        ws[f"I{main_row}"].font = time_font
        ws[f"I{main_row}"].alignment = vertical_up

        ws[f"J{main_row}"] = leg2.trip.trip_no
        ws[f"J{main_row}"].font = trip_font
        ws[f"J{main_row}"].alignment = center
        ws[f"J{main_row}"].border = box_border
        ws[f"J{main_row}"].fill = PatternFill("solid", fgColor=WHITE)
        ws[f"J{bar_row}"].fill = PatternFill("solid", fgColor=color2)
        ws[f"J{bar_row}"].border = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws[f"K{main_row}"] = leg2.trip.arr_time
        ws[f"K{main_row}"].number_format = "hh:mm"
        ws[f"K{main_row}"].font = time_font
        ws[f"K{main_row}"].alignment = vertical_up

        if leg2.role == Role.PASSENGER:
            for col_letter in ("I", "J", "K"):
                ws[f"{col_letter}{main_row}"].font = Font(
                    name=FONT_NAME, bold=True, size=(12 if col_letter == "J" else 8),
                    color="666666", italic=True,
                )

    n_formula = f"=M{main_row}-C{main_row}"
    ws[f"N{main_row}"] = n_formula
    ws[f"N{main_row}"].number_format = "[h]:mm"
    ws[f"N{main_row}"].font = normal_font
    ws[f"N{main_row}"].border = plain_border
    ws[f"N{main_row}"].alignment = center
    _merge_v(ws, 14, main_row, bar_row)

    ws[f"O{main_row}"] = f'=IF(N{main_row}>TIME(8,0,0),"OT","")'
    ws[f"O{main_row}"].font = bold_font
    ws[f"O{main_row}"].border = plain_border
    ws[f"O{main_row}"].alignment = center
    ws[f"O{main_row}"].fill = PatternFill("solid", fgColor=OT_FILL if duty.overtime else "F2F2F2")
    _merge_v(ws, 15, main_row, bar_row)

    for col_letter in list("ABCDEFGHIJKLMNO"):
        top_c = ws[f"{col_letter}{main_row}"]
        bot_c = ws[f"{col_letter}{bar_row}"]
        top_c.border = Border(left=top_c.border.left, right=top_c.border.right,
                               top=medium, bottom=top_c.border.bottom)
        bot_c.border = Border(left=bot_c.border.left, right=bot_c.border.right,
                               top=bot_c.border.top, bottom=medium)

    ws.row_dimensions[main_row].height = 40
    ws.row_dimensions[bar_row].height = 6


def _patch_cached_formula_values(path: str, sheet_file: str, patches: dict[str, tuple[str, bool]]) -> None:
    """See order_b_engine/README.md: openpyxl writes formulas with no cached value, so a
    non-recalculating viewer shows 0/blank until real Excel recalculates. This bakes in the
    known-correct cached value so the file displays right everywhere, while formulas stay live."""
    with zipfile.ZipFile(path, "r") as zin:
        items = zin.infolist()
        data = {i.filename: zin.read(i.filename) for i in items}

    xml = data[sheet_file].decode("utf-8")
    for cell_ref, (value_str, is_str) in patches.items():
        pattern = re.compile(r'(<c r="' + cell_ref + r'"[^>]*>)(<f>.*?</f>)<v\s*/>(</c>)', re.DOTALL)

        def repl(m, value_str=value_str, is_str=is_str):
            open_tag = m.group(1)
            if is_str and " t=" not in open_tag:
                open_tag = open_tag[:-1] + ' t="str">'
            return f"{open_tag}{m.group(2)}<v>{value_str}</v>{m.group(3)}"

        xml, n = pattern.subn(repl, xml)
        if n != 1:
            raise RuntimeError(f"cached-value patch failed for {cell_ref} (matched {n})")
    data[sheet_file] = xml.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in items:
            zout.writestr(i, data[i.filename])


def write_task_schedule(duties: list[Duty], path: str, title: str = "Driver Task Schedule") -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Driver Tasks"

    ws["A1"] = title
    ws["A1"].font = title_font
    ws.merge_cells("A1:O1")

    row = _write_legend(ws, 4)
    header_row = row
    _write_header_row(ws, header_row)

    formula_patches: dict[str, tuple[str, bool]] = {}
    main_row = header_row + 1
    for duty in duties:
        _write_duty_row(ws, duty, main_row)
        duty_minutes = duty.total_duty_minutes()
        formula_patches[f"N{main_row}"] = (repr(duty_minutes / (24 * 60)), False)
        formula_patches[f"O{main_row}"] = ("OT" if duty.overtime else "", True)
        main_row += 2

    widths = {
        "A": 10, "B": 30, "C": 8, "D": 10,
        "E": 6, "F": 11, "G": 6, "H": 5,
        "I": 6, "J": 11, "K": 6, "L": 10, "M": 8, "N": 9, "O": 7,
    }
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w
    ws.freeze_panes = f"A{header_row + 1}"

    wb.save(path)
    if formula_patches:
        _patch_cached_formula_values(path, "xl/worksheets/sheet1.xml", formula_patches)
